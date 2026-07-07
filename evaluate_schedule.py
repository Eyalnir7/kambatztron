"""Evaluate an existing schedule from an xlsx file."""
import argparse
from pathlib import Path
from readers.cadet_reader import CadetCSVReader
from readers.job_reader import JobJSONReader
from output.evaluator import Evaluator
from scheduling.schedule import Schedule
from domain.job import Shift
from domain.time_slot import TimeSlot
from openpyxl import load_workbook


def find_cadet_by_name(cadets, name):
    """Find a cadet by name."""
    return next((c for c in cadets if c.name == name), None)


def evaluate_xlsx(cadets_path: str, jobs_path: str, xlsx_path: str, output_path: str) -> None:
    """Evaluate a schedule from an xlsx file created in matrix format.
    
    The xlsx file is expected to have sheets named by job_type, with:
    - First column: job names
    - Other columns: time slots
    - Cell values: cadet names (or empty for unassigned)
    """
    # Read input files
    cadets = CadetCSVReader.read(cadets_path)
    jobs = JobJSONReader.read(jobs_path)
    
    # Create lookup maps
    cadet_map = {c.personal_number: c for c in cadets}
    job_map = {j.name: j for j in jobs}
    
    # Load xlsx file
    wb = load_workbook(xlsx_path)
    
    assignments = []
    
    # Process each sheet (excluding legend sheet)
    for sheet_name in wb.sheetnames:
        if sheet_name == "Legend":
            continue
            
        ws = wb[sheet_name]
        
        # First row contains time slots (skip first column which has job names)
        time_slots = []
        for col_idx in range(2, ws.max_column + 1):
            ts_str = ws.cell(row=1, column=col_idx).value
            if ts_str:
                try:
                    ts = TimeSlot.from_string(str(ts_str))
                    time_slots.append((col_idx, ts))
                except Exception as e:
                    print(f"Warning: Could not parse time slot '{ts_str}': {e}")
        
        # Process each job (rows)
        for row_idx in range(2, ws.max_row + 1):
            job_name = ws.cell(row=row_idx, column=1).value
            if not job_name or job_name not in job_map:
                continue
            
            job = job_map[job_name]
            
            # Process each time slot for this job
            for col_idx, time_slot in time_slots:
                cadet_name = ws.cell(row=row_idx, column=col_idx).value
                
                # Get assigned cadet
                assigned_cadet = None
                if cadet_name:
                    assigned_cadet = find_cadet_by_name(cadets, str(cadet_name))
                
                # Get difficulty score for this time slot
                difficulty = job.difficulty_by_slot.get(time_slot, 5.0)
                
                # Create shift
                shift = Shift(
                    job=job,
                    time_slot=time_slot,
                    assigned_cadet=assigned_cadet,
                    difficulty=float(difficulty)
                )
                assignments.append(shift)
    
    # Create schedule
    schedule = Schedule(assignments=assignments)
    
    # Run evaluation
    evaluator = Evaluator()
    evaluator.evaluate(schedule, cadets, output_path)
    print(f"Evaluation complete! Results saved to {output_path}/evaluation/")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Evaluate an existing schedule from xlsx")
    parser.add_argument("--cadets", required=True, help="Path to cadets CSV file")
    parser.add_argument("--jobs", required=True, help="Path to jobs JSON file")
    parser.add_argument("--xlsx", required=True, help="Path to the xlsx schedule file")
    parser.add_argument("--output", required=True, help="Output directory for evaluation results")
    
    args = parser.parse_args()
    evaluate_xlsx(args.cadets, args.jobs, args.xlsx, args.output)
